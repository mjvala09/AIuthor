import os
import json
import logging
import re
from typing import List, Dict, Any, Tuple, Optional
from app.config import settings
from app.models.schemas import BookOutline

# ReportLab Imports
from reportlab.lib.pagesizes import letter
from reportlab.lib import colors
from reportlab.lib.units import inch
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, PageBreak, Table, TableStyle, KeepTogether, Flowable
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.pdfgen import canvas

# Docx Imports
try:
    import docx
    from docx.shared import Pt, Inches, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.oxml import OxmlElement, parse_xml
    from docx.oxml.ns import nsdecls, qn
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

logger = logging.getLogger("doc_gen")

# Helper to convert to roman numerals
def to_roman(n: int) -> str:
    if n <= 0:
        return ""
    val = [1000, 900, 500, 400, 100, 90, 50, 40, 10, 9, 5, 4, 1]
    syb = ["M", "CM", "D", "CD", "C", "XC", "L", "XL", "X", "IX", "V", "IV", "I"]
    roman_num = ''
    i = 0
    while n > 0:
        for _ in range(n // val[i]):
            roman_num += syb[i]
            n -= val[i]
        i += 1
    return roman_num.lower()

# Markdown to HTML compiler for ReportLab Paragraphs
def md_to_html(text: str) -> str:
    if not text:
        return ""
    # 1. Escape basic XML entities
    text = text.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
    
    # 2. Match bold **text** or __text__
    text = re.sub(r"\*\*(.*?)\*\*", r"<b>\1</b>", text)
    text = re.sub(r"__(.*?)__", r"<b>\1</b>", text)
    
    # 3. Match italic *text* or _text_
    text = re.sub(r"\*(.*?)\*", r"<i>\1</i>", text)
    text = re.sub(r"_(.*?)_", r"<i>\1</i>", text)
    
    # Replace newlines with <br/>
    text = text.replace("\n", "<br/>")
    return text

# Custom Flowable to record page numbers and chapter titles during PDF pass
class PageTracker(Flowable):
    def __init__(self, key, title, callback_dict):
        super().__init__()
        self.key = key
        self.title = title
        self.callback_dict = callback_dict
        self.width = 0
        self.height = 0

    def draw(self):
        self.callback_dict[self.key] = self.canv._pageNumber
        self.canv.current_chapter_title = self.title
        if self.key == "introduction":
            self.canv.body_start_page = self.canv._pageNumber
        
        # Register the page bookmark destination for clickable internal links
        self.canv.bookmarkPage(self.key)
        
        # Register a PDF outline side-panel entry
        self.canv.addOutlineEntry(self.title, self.key, level=0)

# Helper functions for python-docx bookmarks and links
def add_bookmark(paragraph, name, bookmark_id):
    p = paragraph._p
    bookmark_start = OxmlElement('w:bookmarkStart')
    bookmark_start.set(qn('w:id'), str(bookmark_id))
    bookmark_start.set(qn('w:name'), name)
    p.append(bookmark_start)
    
    bookmark_end = OxmlElement('w:bookmarkEnd')
    bookmark_end.set(qn('w:id'), str(bookmark_id))
    p.append(bookmark_end)

def add_hyperlink_to_bookmark(paragraph, bookmark_name, text, color="0000FF", underline=True):
    p = paragraph._p
    hyperlink = OxmlElement('w:hyperlink')
    hyperlink.set(qn('w:anchor'), bookmark_name)
    hyperlink.set(qn('w:history'), '1')
    
    new_run = OxmlElement('w:r')
    rPr = OxmlElement('w:rPr')
    
    if color:
        c = OxmlElement('w:color')
        c.set(qn('w:val'), color)
        rPr.append(c)
    if underline:
        u = OxmlElement('w:u')
        u.set(qn('w:val'), 'single')
        rPr.append(u)
        
    new_run.append(rPr)
    
    text_node = OxmlElement('w:t')
    text_node.text = text
    new_run.append(text_node)
    
    hyperlink.append(new_run)
    p.append(hyperlink)
    return docx.text.run.Run(new_run, paragraph)

def add_markdown_paragraph(doc, text, style=None):
    """Parses markdown bold (**) and italic (* or _) and adds a formatted paragraph to the document."""
    p = doc.add_paragraph(style=style)
    p.style.font.name = 'Arial'
    
    tokens = re.split(r'(\*\*\*.*?\*\*\*|\*\*.*?\*\*|\*.*?\*|___.*?___|__.*?__|__.*?__|_.*?_)', text)
    
    for token in tokens:
        if not token:
            continue
        
        bold = False
        italic = False
        content = token
        
        if token.startswith('***') and token.endswith('***'):
            bold = True
            italic = True
            content = token[3:-3]
        elif token.startswith('___') and token.endswith('___'):
            bold = True
            italic = True
            content = token[3:-3]
        elif token.startswith('**') and token.endswith('**'):
            bold = True
            content = token[2:-2]
        elif token.startswith('__') and token.endswith('__'):
            bold = True
            content = token[2:-2]
        elif token.startswith('*') and token.endswith('*'):
            italic = True
            content = token[1:-1]
        elif token.startswith('_') and token.endswith('_'):
            italic = True
            content = token[1:-1]
            
        run = p.add_run(content)
        run.font.name = 'Arial'
        if bold:
            run.bold = True
        if italic:
            run.italic = True
            
    return p


# Custom Canvas for dynamic footers & headers
class BookCanvas(canvas.Canvas):
    def __init__(self, *args, **kwargs):
        super().__init__(*args, **kwargs)
        self.pages = []
        self.chapter_targets = {}  # Map of target names -> absolute page number
        self.body_start_page = 0
        self.current_chapter_title = ""

    def showPage(self):
        # Record state for the second-pass draw
        self.pages.append({
            "dict": self.__dict__.copy(),
            "chapter_title": self.current_chapter_title,
            "body_start_page": self.body_start_page
        })
        self._startPage()

    def save(self):
        page_count = len(self.pages)
        # Find when the body starts (typically when page template switches to Body)
        # Or we can scan page templates or hardcode based on first pass
        first_body_page = 0
        for idx, p in enumerate(self.pages):
            if "body" in str(p["dict"].get("_templateId", "")).lower() or p["dict"].get("body_start_page", 0) > 0:
                first_body_page = idx + 1
                break
                
        if first_body_page == 0:
            first_body_page = 8  # fallback approximation
            
        for idx, page in enumerate(self.pages):
            self.__dict__.update(page["dict"])
            page_num = idx + 1
            
            # Skip headers/footers on cover pages
            template_id = str(self.__dict__.get("_templateId", "")).lower()
            is_cover = "cover" in template_id or page_num == 1 or page_num == 2
            
            if not is_cover:
                # Draw Header
                self.saveState()
                self.setFont("Helvetica-Oblique", 9)
                self.setFillColor(colors.HexColor("#4B5563")) # Gray-600
                
                # Header text
                header_text = page["chapter_title"] if page["chapter_title"] else "AIuthor Book Generator"
                if page_num % 2 == 0:
                    self.drawString(54, 750, header_text)
                else:
                    self.drawRightString(612 - 54, 750, header_text)
                    
                # Header rule
                self.setStrokeColor(colors.HexColor("#E5E7EB"))
                self.setLineWidth(0.5)
                self.line(54, 742, 612 - 54, 742)
                
                # Draw Footer
                self.line(54, 50, 612 - 54, 50)
                self.setFont("Helvetica", 9)
                
                if page_num < first_body_page:
                    # Roman numerals for front matter
                    roman_page = to_roman(page_num - 2) if page_num > 2 else to_roman(page_num)
                    self.drawCentredString(306, 35, roman_page)
                else:
                    # Arabic starting from 1 for the introduction/body
                    body_page = page_num - first_body_page + 1
                    self.drawCentredString(306, 35, str(body_page))
                    
                self.restoreState()
                
            super().showPage()
        super().save()

class DocGenService:
    def _sanitize_matter(self, matter: Any) -> Dict[str, str]:
        if not isinstance(matter, dict):
            matter = {}
        sanitized = {}
        for k, v in matter.items():
            if isinstance(v, dict):
                lines = []
                for key, val in v.items():
                    if isinstance(val, (dict, list)):
                        val = json.dumps(val)
                    lines.append(f"{key.replace('_', ' ').title()}: {val}")
                sanitized[k] = "\n".join(lines)
            elif isinstance(v, list):
                lines = []
                for item in v:
                    if isinstance(item, (dict, list)):
                        item = json.dumps(item)
                    lines.append(str(item))
                sanitized[k] = "\n".join(lines)
            elif v is None:
                sanitized[k] = ""
            else:
                sanitized[k] = str(v)
        return sanitized

    def generate_documents(
        self,
        run_id: str,
        outline: BookOutline,
        chapters: List[Dict[str, Any]],
        matter: Dict[str, Any]
    ) -> Tuple[str, str]:
        """Generates beautifully styled PDF and DOCX files for the book."""
        # Sanitize matter to ensure all values are strings
        matter = self._sanitize_matter(matter)
        pdf_path = os.path.join(settings.BOOKS_DIR, f"{run_id}_book.pdf")
        docx_path = os.path.join(settings.BOOKS_DIR, f"{run_id}_book.docx")
        
        # 1. Generate PDF
        self._generate_pdf(pdf_path, outline, chapters, matter)
        
        # 2. Generate DOCX
        if DOCX_AVAILABLE:
            self._generate_docx(docx_path, outline, chapters, matter)
        else:
            # Create a mock docx if library is missing
            with open(docx_path, "w", encoding="utf-8") as f:
                f.write(f"DOCX placeholder for {outline.title}")
                
        return pdf_path, docx_path

    def _generate_pdf(
        self,
        path: str,
        outline: BookOutline,
        chapters: List[Dict[str, Any]],
        matter: Dict[str, str]
    ):
        # We need two passes to resolve exact page numbers for the TOC!
        # First Pass: build without TOC numbers to record heading page numbers
        chapter_pages = self._run_pdf_pass(None, outline, chapters, matter)
        
        # Second Pass: build with real page numbers in the TOC!
        self._run_pdf_pass(path, outline, chapters, matter, chapter_pages)

    def _run_pdf_pass(
        self,
        path: Optional[str],
        outline: BookOutline,
        chapters: List[Dict[str, Any]],
        matter: Dict[str, str],
        toc_pages: Optional[Dict[str, int]] = None
    ) -> Dict[str, int]:
        
        # If path is None, we write to a temporary buffer to get page numbers
        target_path = path or os.path.join(settings.MEMORIES_DIR, "temp_pass.pdf")
        
        doc = SimpleDocTemplate(
            target_path,
            pagesize=letter,
            leftMargin=54,
            rightMargin=54,
            topMargin=72,
            bottomMargin=72
        )
        
        styles = getSampleStyleSheet()
        
        # Custom Styles
        title_style = ParagraphStyle(
            'BookTitle',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=32,
            leading=38,
            textColor=colors.HexColor("#111827"), # Dark gray
            alignment=1, # Centered
            spaceAfter=15
        )
        
        subtitle_style = ParagraphStyle(
            'BookSubtitle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=18,
            leading=24,
            textColor=colors.HexColor("#4B5563"),
            alignment=1,
            spaceAfter=40
        )
        
        author_style = ParagraphStyle(
            'BookAuthor',
            parent=styles['Normal'],
            fontName='Helvetica-Bold',
            fontSize=14,
            leading=18,
            textColor=colors.HexColor("#1F2937"),
            alignment=1,
            spaceAfter=100
        )
        
        h1_style = ParagraphStyle(
            'BookH1',
            parent=styles['Heading1'],
            fontName='Helvetica-Bold',
            fontSize=22,
            leading=26,
            textColor=colors.HexColor("#1F2937"),
            spaceBefore=20,
            spaceAfter=15,
            keepWithNext=True
        )
        
        body_style = ParagraphStyle(
            'BookBody',
            parent=styles['BodyText'],
            fontName='Helvetica',
            fontSize=11,
            leading=16,
            textColor=colors.HexColor("#374151"),
            spaceAfter=10
        )
        
        toc_style = ParagraphStyle(
            'TOCStyle',
            parent=styles['Normal'],
            fontName='Helvetica',
            fontSize=11,
            leading=16,
            textColor=colors.HexColor("#1F2937")
        )
        
        story = []
        
        # --- 1. COVER PAGE ---
        story.append(Spacer(1, 150))
        story.append(Paragraph(outline.title, title_style))
        if outline.subtitle:
            story.append(Paragraph(outline.subtitle, subtitle_style))
        story.append(Spacer(1, 50))
        story.append(Paragraph(f"By {outline.author or 'AIuthor'}", author_style))
        story.append(PageBreak())
        
        # --- 2. COPYRIGHT PAGE ---
        story.append(Spacer(1, 40))
        story.append(Paragraph("<b>Copyright</b>", h1_style))
        story.append(Paragraph(matter.get("copyright", "© 2026 AIuthor.").replace("\n", "<br/>"), body_style))
        story.append(PageBreak())
        
        # --- 3. DEDICATION ---
        story.append(Spacer(1, 100))
        story.append(Paragraph(f"<i>{matter.get('dedication', 'Dedicated to readers.')}</i>", subtitle_style))
        story.append(PageBreak())
        
        # --- 4. EPIGRAPH ---
        story.append(Spacer(1, 120))
        story.append(Paragraph(f"<i>{matter.get('epigraph', 'Knowledge is power.')}</i>", subtitle_style))
        story.append(PageBreak())
        
        # --- 5. TABLE OF CONTENTS ---
        story.append(Spacer(1, 20))
        story.append(Paragraph("Table of Contents", h1_style))
        story.append(Spacer(1, 10))
        
        # Generate static/dynamic rows for TOC
        toc_data = []
        
        # Add Front Matter entries
        front_matters = [
            ("Foreword", "foreword"),
            ("Preface", "preface"),
            ("Acknowledgments", "acknowledgments"),
            ("Introduction", "introduction")
        ]
        
        for name, key in front_matters:
            p_num = toc_pages.get(key, 0) if toc_pages else 0
            p_num_str = to_roman(p_num - 2) if p_num > 2 else to_roman(p_num)
            if not p_num_str:
                p_num_str = "..."
            toc_data.append([
                Paragraph(f"<a href=\"#{key}\"><b>{name}</b></a>", toc_style),
                Paragraph(". " * 30, toc_style),
                Paragraph(f"<a href=\"#{key}\">{p_num_str}</a>", toc_style)
            ])
            
        # Add Chapter entries
        for ch in chapters:
            key = f"chapter_{ch['chapter_number']}"
            intro_abs_page = toc_pages.get("introduction", 8) if toc_pages else 8
            p_num = toc_pages.get(key, 0) if toc_pages else 0
            p_num_str = str(p_num - intro_abs_page + 1) if p_num > 0 else "..."
            
            toc_data.append([
                Paragraph(f"<a href=\"#{key}\">Chapter {ch['chapter_number']}: {ch['title']}</a>", toc_style),
                Paragraph(". " * 30, toc_style),
                Paragraph(f"<a href=\"#{key}\">{p_num_str}</a>", toc_style)
            ])
            
        # Add Back Matter entries
        back_matters = [
            ("Afterword", "afterword"),
            ("Appendix", "appendix"),
            ("Glossary", "glossary"),
            ("References", "references"),
            ("About the Author", "about_author")
        ]
        for name, key in back_matters:
            intro_abs_page = toc_pages.get("introduction", 8) if toc_pages else 8
            p_num = toc_pages.get(key, 0) if toc_pages else 0
            p_num_str = str(p_num - intro_abs_page + 1) if p_num > 0 else "..."
            toc_data.append([
                Paragraph(f"<a href=\"#{key}\"><b>{name}</b></a>", toc_style),
                Paragraph(". " * 30, toc_style),
                Paragraph(f"<a href=\"#{key}\">{p_num_str}</a>", toc_style)
            ])
            
        t = Table(toc_data, colWidths=[200, 250, 50])
        t.setStyle(TableStyle([
            ('VALIGN', (0,0), (-1,-1), 'BOTTOM'),
            ('BOTTOMPADDING', (0,0), (-1,-1), 2),
            ('TOPPADDING', (0,0), (-1,-1), 2)
        ]))
        story.append(t)
        story.append(PageBreak())
        
        # Helper to append sections and track their page numbers
        pages_dict = {}
        
        # --- 6. FRONT MATTER ---
        for name, key in front_matters:
            story.append(Spacer(1, 20))
            story.append(PageTracker(key, name, pages_dict))
            heading_p = Paragraph(f"<a name=\"{key}\"/>{name}", h1_style)
            story.append(heading_p)
            story.append(Paragraph(md_to_html(matter.get(key, "")), body_style))
            story.append(PageBreak())
            
        # --- 7. CHAPTERS ---
        for ch in chapters:
            key = f"chapter_{ch['chapter_number']}"
            ch_title = f"Chapter {ch['chapter_number']}: {ch['title']}"
            story.append(Spacer(1, 20))
            story.append(PageTracker(key, ch_title, pages_dict))
            story.append(Paragraph(f"<a name=\"{key}\"/>{ch_title}", h1_style))
            
            # Write text
            # Format markdown paragraphs
            paragraphs = ch["final_text"].split("\n\n")
            for p in paragraphs:
                if p.strip().startswith("###"):
                    story.append(Paragraph(md_to_html(p.replace("###", "").strip()), styles["Heading3"]))
                elif p.strip().startswith("##"):
                    story.append(Paragraph(md_to_html(p.replace("##", "").strip()), styles["Heading2"]))
                elif p.strip():
                    story.append(Paragraph(md_to_html(p.strip()), body_style))
                    
            story.append(PageBreak())
            
        # --- 8. BACK MATTER ---
        for name, key in back_matters:
            story.append(Spacer(1, 20))
            story.append(PageTracker(key, name, pages_dict))
            story.append(Paragraph(f"<a name=\"{key}\"/>{name}", h1_style))
            
            paragraphs = matter.get(key, "").split("\n\n")
            for p in paragraphs:
                if p.strip():
                    story.append(Paragraph(md_to_html(p.strip()), body_style))
            story.append(PageBreak())
            
        # Build Document using custom Canvas
        doc.build(story, canvasmaker=BookCanvas)
        
        # Clean up temp file if first pass
        if not path and os.path.exists(target_path):
            try:
                os.remove(target_path)
            except Exception:
                pass
                
        # Return page number catalog
        if not pages_dict:
            # Fallback estimation for first pass
            pages_dict = {
                "foreword": 3,
                "preface": 4,
                "acknowledgments": 5,
                "introduction": 6,
                "afterword": 20,
                "appendix": 21,
                "glossary": 22,
                "references": 23,
                "about_author": 24
            }
            curr_p = 7
            for ch in chapters:
                pages_dict[f"chapter_{ch['chapter_number']}"] = curr_p
                curr_p += 2  # assume 2 pages per chapter
                
        return pages_dict

    def _generate_docx(
        self,
        path: str,
        outline: BookOutline,
        chapters: List[Dict[str, Any]],
        matter: Dict[str, str]
    ):
        doc = docx.Document()
        
        # Set standard margins (1 inch)
        for section in doc.sections:
            section.top_margin = Inches(1)
            section.bottom_margin = Inches(1)
            section.left_margin = Inches(1)
            section.right_margin = Inches(1)
            
        # Title Style
        title_p = doc.add_paragraph()
        title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        title_run = title_p.add_run(outline.title)
        title_run.font.name = 'Arial'
        title_run.font.size = Pt(28)
        title_run.bold = True
        
        if outline.subtitle:
            sub_p = doc.add_paragraph()
            sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            sub_run = sub_p.add_run(outline.subtitle)
            sub_run.font.name = 'Arial'
            sub_run.font.size = Pt(16)
            sub_run.font.color.rgb = RGBColor(100, 100, 100)
            
        doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_p = doc.add_paragraph()
        author_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        author_run = author_p.add_run(f"By {outline.author or 'AIuthor'}")
        author_run.font.name = 'Arial'
        author_run.font.size = Pt(14)
        author_run.bold = True
        
        doc.add_page_break()
        
        # Copyright
        h = doc.add_heading("Copyright", level=1)
        h.runs[0].font.name = 'Arial'
        p = doc.add_paragraph(matter.get("copyright", ""))
        p.style.font.name = 'Arial'
        doc.add_page_break()
        
        # Dedication
        doc.add_paragraph()
        p = doc.add_paragraph(matter.get("dedication", ""))
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        if p.runs:
            p.runs[0].italic = True
        p.style.font.name = 'Arial'
        doc.add_page_break()
        
        # Table of Contents
        doc.add_heading("Table of Contents", level=1)
        
        bookmark_id = 1
        
        # Front Matter TOC
        for item in ["Foreword", "Preface", "Acknowledgments", "Introduction"]:
            key = item.lower()
            p = doc.add_paragraph()
            add_hyperlink_to_bookmark(p, key, item, color="1F2937", underline=False)
            p.add_run(" " + "." * (80 - len(item)) + " ")
            add_hyperlink_to_bookmark(p, key, "[Go]", color="2563EB", underline=True)
            
        # Chapters TOC
        for ch in chapters:
            key = f"chapter_{ch['chapter_number']}"
            title_text = f"Chapter {ch['chapter_number']}: {ch['title']}"
            p = doc.add_paragraph()
            add_hyperlink_to_bookmark(p, key, title_text, color="1F2937", underline=False)
            p.add_run(" " + "." * (80 - len(title_text)) + " ")
            add_hyperlink_to_bookmark(p, key, "[Go]", color="2563EB", underline=True)
            
        # Back Matter TOC
        back_matters = [
            ("Afterword", "afterword"),
            ("Appendix", "appendix"),
            ("Glossary", "glossary"),
            ("References", "references"),
            ("About the Author", "about_author")
        ]
        for name, key in back_matters:
            p = doc.add_paragraph()
            add_hyperlink_to_bookmark(p, key, name, color="1F2937", underline=False)
            p.add_run(" " + "." * (80 - len(name)) + " ")
            add_hyperlink_to_bookmark(p, key, "[Go]", color="2563EB", underline=True)
            
        doc.add_page_break()
        
        # Front Matter Sections
        for item in ["foreword", "preface", "acknowledgments", "introduction"]:
            title = item.capitalize() if item != "about_author" else "About the Author"
            h = doc.add_heading(title, level=1)
            h.runs[0].font.name = 'Arial'
            add_bookmark(h, item, bookmark_id)
            bookmark_id += 1
            
            # Format markdown paragraphs
            paragraphs = matter.get(item, "").split("\n\n")
            for p_text in paragraphs:
                if p_text.strip():
                    add_markdown_paragraph(doc, p_text.strip())
            doc.add_page_break()
            
        # Chapters Sections
        for ch in chapters:
            ch_title = f"Chapter {ch['chapter_number']}: {ch['title']}"
            h = doc.add_heading(ch_title, level=1)
            h.runs[0].font.name = 'Arial'
            add_bookmark(h, f"chapter_{ch['chapter_number']}", bookmark_id)
            bookmark_id += 1
            
            paragraphs = ch["final_text"].split("\n\n")
            for p_text in paragraphs:
                if p_text.strip().startswith("###"):
                    sub_title = p_text.replace("###", "").strip()
                    add_markdown_paragraph(doc, sub_title, style='Heading 3')
                elif p_text.strip().startswith("##"):
                    sub_title = p_text.replace("##", "").strip()
                    add_markdown_paragraph(doc, sub_title, style='Heading 2')
                elif p_text.strip():
                    add_markdown_paragraph(doc, p_text.strip())
            doc.add_page_break()
            
        # Back Matter Sections
        for name, key in back_matters:
            h = doc.add_heading(name, level=1)
            h.runs[0].font.name = 'Arial'
            add_bookmark(h, key, bookmark_id)
            bookmark_id += 1
            
            paragraphs = matter.get(key, "").split("\n\n")
            for p_text in paragraphs:
                if p_text.strip():
                    add_markdown_paragraph(doc, p_text.strip())
            doc.add_page_break()
            
        doc.save(path)

doc_gen_service = DocGenService()
